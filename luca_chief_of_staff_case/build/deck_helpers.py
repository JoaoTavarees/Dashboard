"""Shared slide layout helpers (python-pptx) so both decks look like one family."""
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(0x1F, 0x2A, 0x44)
TEAL = RGBColor(0x0F, 0x9D, 0x8A)
TEAL_LT = RGBColor(0xE6, 0xF6, 0xF3)
CORAL = RGBColor(0xE8, 0x6A, 0x4A)
CORAL_LT = RGBColor(0xFD, 0xEE, 0xE9)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)
AMBER_LT = RGBColor(0xFE, 0xF6, 0xE4)
GREY = RGBColor(0x6B, 0x72, 0x80)
GREY_LT = RGBColor(0xF3, 0xF4, 0xF6)
LINE = RGBColor(0xD1, 0xD5, 0xDB)
INK = RGBColor(0x11, 0x18, 0x27)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"


def new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    return s


def _runs(p, text, size, color, bold=False, italic=False):
    """Supports **bold** inline markup."""
    for k, part in enumerate(re.split(r"(\*\*.+?\*\*)", text)):
        if not part:
            continue
        r = p.add_run()
        is_b = part.startswith("**") and part.endswith("**")
        r.text = part[2:-2] if is_b else part
        f = r.font
        f.name, f.size, f.bold, f.italic = FONT, Pt(size), bold or is_b, italic
        f.color.rgb = color


def text(slide, x, y, w, h, paras, size=12, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         spacing=4, margin=0.05):
    """paras: list of str or dict(t=..., size=, color=, bold=, bullet=bool, italic=, space=)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    for i, para in enumerate(paras):
        if isinstance(para, str):
            para = dict(t=para)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = para.get("align", align)
        p.space_after = Pt(para.get("space", spacing))
        t = para["t"]
        if para.get("bullet"):
            t = "•  " + t
        _runs(p, t, para.get("size", size), para.get("color", color), para.get("bold", False),
              para.get("italic", False))
    return tb


def base_slide(prs, kicker, title, page, source=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, 0.12, 7.5, TEAL)
    text(s, 0.5, 0.28, 12.3, 0.3, [dict(t=kicker.upper(), size=10.5, bold=True, color=TEAL)])
    text(s, 0.5, 0.52, 12.3, 0.9, [dict(t=title, size=21, bold=True, color=NAVY)], spacing=0)
    s.shapes.add_connector(1, Inches(0.5), Inches(1.45), Inches(12.83), Inches(1.45)).line.color.rgb = LINE
    foot = "Luca · Chief of Staff hiring case · all data fictional"
    if source:
        foot += "   |   " + source
    text(s, 0.5, 7.08, 11.5, 0.3, [dict(t=foot, size=8.5, color=GREY)])
    text(s, 12.2, 7.08, 0.65, 0.3, [dict(t=str(page), size=8.5, color=GREY, align=PP_ALIGN.RIGHT)])
    return s


def card(slide, x, y, w, h, label, value, sub=None, fill=TEAL_LT, accent=TEAL, value_size=24):
    rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(slide, x, y + 0.12, 0.06, h - 0.24, accent)
    paras = [dict(t=label.upper(), size=9, bold=True, color=GREY, space=2),
             dict(t=value, size=value_size, bold=True, color=NAVY, space=2)]
    if sub:
        paras.append(dict(t=sub, size=9.5, color=INK, space=0))
    text(slide, x + 0.18, y + 0.08, w - 0.28, h - 0.12, paras)


def panel(slide, x, y, w, h, heading, paras, fill=GREY_LT, accent=NAVY, size=11, head_size=12.5):
    rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(slide, x + 0.15, y + 0.1, w - 0.3, 0.4, [dict(t=heading, size=head_size, bold=True, color=accent)])
    text(slide, x + 0.15, y + 0.5, w - 0.3, h - 0.55, paras, size=size)


def table(slide, x, y, w, h, rows, col_w=None, size=10, header_fill=NAVY, zebra=True, bold_first_col=False,
          align_right_from=None, row_h=None):
    nr, nc = len(rows), len(rows[0])
    gt = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = gt.table
    tbl.first_row = True
    if col_w:
        for j, cw in enumerate(col_w):
            tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        if row_h:
            tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            fill, color, bold = None, INK, False
            if isinstance(val, dict):
                fill, color, bold = val.get("fill"), val.get("color", INK), val.get("bold", False)
                val = val["t"]
            if i == 0:
                fill, color, bold = header_fill, WHITE, True
            elif fill is None:
                fill = GREY_LT if (zebra and i % 2 == 0) else WHITE
            if bold_first_col and j == 0 and i > 0:
                bold = True
            _runs(p, str(val), size, color, bold)
            if align_right_from is not None and j >= align_right_from:
                p.alignment = PP_ALIGN.RIGHT
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill
    return tbl


def pill(slide, x, y, w, h, label, fill=TEAL, color=WHITE, size=10):
    rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(slide, x, y, w, h, [dict(t=label, size=size, bold=True, color=color, align=PP_ALIGN.CENTER)],
         anchor=MSO_ANCHOR.MIDDLE, margin=0.02)


def style_chart(chart, legend=True, size=9):
    chart.font.name = FONT
    chart.font.size = Pt(size)
    chart.font.color.rgb = INK
    chart.has_legend = legend
    if legend:
        from pptx.enum.chart import XL_LEGEND_POSITION
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(size)
