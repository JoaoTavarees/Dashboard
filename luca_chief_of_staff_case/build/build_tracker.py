"""Builds the content-pipeline master tracker template, pre-loaded with the Week 1 snapshot."""
import datetime as dt
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parents[1] / "deliverables" / "Luca_Content_Pipeline_Master_Tracker.xlsx"
random.seed(7)

NAVY, TEAL = "1F2A44", "0F9D8A"
HDR = PatternFill("solid", fgColor=NAVY)
KEY = PatternFill("solid", fgColor=TEAL)
IN = PatternFill("solid", fgColor="FFF7D6")
F = lambda **k: Font(name="Arial", size=10, **k)

STAGES = ["1. Brief & Content", "2. Script", "3. Visual Assets", "4. Motion", "5. Final Video", "Done"]
SLA = {"1. Brief & Content": 1, "2. Script": 2, "3. Visual Assets": 1, "4. Motion": 2, "5. Final Video": 2}
OWNERS = {"1. Brief & Content": ["Ana (Content)", "Luis (Content)"],
          "2. Script": ["Sofía (Script)", "Diego (Script)", "Mariana (Script)"],
          "3. Visual Assets": ["Carla (Design)", "Pablo (Design)"],
          "4. Motion": ["Jorge (Motion)", "Valeria (Motion)"],
          "5. Final Video": ["Rafa (Synthesia)", "Elena (Sound)"],
          "Done": ["—"]}
BLOCKERS = ["", "Content misalignment / rework", "Waiting for input from previous stage",
            "Change request after lock", "Audio / video sync", "Capacity / not started", "Tool / technical issue"]
SUBJECTS = ["Math", "Spanish", "Science", "History", "English", "Geography", "Civics"]
GRADES = ["Primaria 1", "Primaria 2", "Primaria 3", "Primaria 4", "Primaria 5", "Primaria 6",
          "Secundaria 1", "Secundaria 2", "Secundaria 3"]
W1 = dt.date(2026, 10, 5)
AS_OF = dt.date(2026, 10, 9)

wb = Workbook()
ws = wb.active
ws.title = "Tracker"
ws["A1"] = "Luca – content production master tracker (300 videos / 10 weeks)"
ws["A1"].font = Font(name="Arial", size=14, bold=True, color=NAVY)
ws["A2"] = "As-of date:"
ws["B2"] = AS_OF
ws["B2"].number_format = "dd-mmm-yy"
ws["B2"].fill, ws["B2"].font = IN, F(color="0000FF", bold=True)
ws["D2"] = ("Teal headers = the 5 key fields. One row per video. The stage owner updates their row every day "
            "before the 15-minute stand-up.")
ws["D2"].font = F(italic=True, color="6B7280")

cols = [("Unit ID", 10, False), ("Subject", 11, False), ("Grade", 13, False), ("Tier (A/B/C)", 9, False),
        ("Stage", 18, False), ("Accountable owner", 18, True), ("Stage start date", 12, False),
        ("Stage due date (SLA)", 13, True), ("Days in stage", 9, True), ("Days late", 8, False),
        ("Status", 11, False), ("Blocker reason", 32, True), ("Blocked since", 12, False),
        ("Rework loops", 9, True), ("Done date", 11, False), ("Target ship week", 9, False)]
for j, (name, w, key) in enumerate(cols, start=1):
    c = ws.cell(row=4, column=j, value=name)
    c.font = F(bold=True, color="FFFFFF")
    c.fill = KEY if key else HDR
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.column_dimensions[c.column_letter].width = w
ws.row_dimensions[4].height = 42

layout = (["Done"] * 5 + ["5. Final Video"] * 3 + ["4. Motion"] * 7 + ["3. Visual Assets"] * 5 + ["2. Script"] * 15)
blocker_for = {"5. Final Video": BLOCKERS[4], "4. Motion": BLOCKERS[3], "3. Visual Assets": BLOCKERS[2],
               "2. Script": BLOCKERS[1]}
N_ROWS = 300
for i in range(N_ROWS):
    r = 5 + i
    ws[f"A{r}"] = f"V-{i + 1:03d}"
    ws[f"P{r}"] = i // 30 + 1
    ws[f"I{r}"] = f'=IF(OR(E{r}="",E{r}="Done",G{r}=""),"",$B$2-G{r})'
    ws[f"J{r}"] = f'=IF(OR(E{r}="",E{r}="Done",H{r}=""),"",MAX(0,$B$2-H{r}))'
    ws[f"H{r}"] = (f'=IF(OR(G{r}="",E{r}="Done"),"",G{r}+IFERROR(VLOOKUP(E{r},Lists!$A$2:$B$6,2,FALSE),0))')
    if i < len(layout):
        st = layout[i]
        ws[f"B{r}"] = SUBJECTS[i % len(SUBJECTS)]
        ws[f"C{r}"] = random.choice(GRADES)
        ws[f"D{r}"] = random.choice(["A", "B", "B", "C"])
        ws[f"E{r}"] = st
        ws[f"F{r}"] = random.choice(OWNERS[st])
        if st == "Done":
            ws[f"K{r}"] = "Done"
            ws[f"O{r}"] = AS_OF - dt.timedelta(days=random.choice([0, 1]))
            ws[f"N{r}"] = random.choice([1, 1, 2])
        else:
            ws[f"G{r}"] = AS_OF - dt.timedelta(days={"2. Script": random.choice([2, 3, 4]),
                                                      "3. Visual Assets": 2, "4. Motion": random.choice([2, 3]),
                                                      "5. Final Video": 1}[st])
            ws[f"K{r}"] = "Blocked"
            ws[f"L{r}"] = blocker_for[st]
            ws[f"M{r}"] = ws[f"G{r}"].value + dt.timedelta(days=1)
            ws[f"N{r}"] = random.choice([1, 2, 2, 3]) if st == "2. Script" else random.choice([0, 1, 1])
    else:
        ws[f"K{r}"] = "Not started"
        ws[f"N{r}"] = 0
    for col in "GHMO":
        ws[f"{col}{r}"].number_format = "dd-mmm"
    for col in "ABCDEFGHIJKLMNOP":
        ws[f"{col}{r}"].font = F()

last = 4 + N_ROWS
dvs = [("E", '"' + ",".join(STAGES) + '"'), ("K", '"Not started,On track,At risk,Blocked,Done"'),
       ("D", '"A,B,C"'), ("L", "=Lists!$D$2:$D$8")]
for col, formula in dvs:
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col}5:{col}{last}")
red = PatternFill("solid", fgColor="FFC7CE")
amber = PatternFill("solid", fgColor="FFEB9C")
green = PatternFill("solid", fgColor="C6EFCE")
ws.conditional_formatting.add(f"K5:K{last}", CellIsRule(operator="equal", formula=['"Blocked"'], fill=red))
ws.conditional_formatting.add(f"K5:K{last}", CellIsRule(operator="equal", formula=['"At risk"'], fill=amber))
ws.conditional_formatting.add(f"K5:K{last}", CellIsRule(operator="equal", formula=['"Done"'], fill=green))
ws.conditional_formatting.add(f"J5:J{last}", CellIsRule(operator="greaterThan", formula=["0"], fill=red))
ws.conditional_formatting.add(f"N5:N{last}", CellIsRule(operator="greaterThanOrEqual", formula=["2"], fill=amber))
ws.conditional_formatting.add(f"I5:I{last}", FormulaRule(formula=[f'AND(ISNUMBER(I5),I5>2)'], fill=amber))
ws.freeze_panes = "B5"
ws.auto_filter.ref = f"A4:P{last}"

ls = wb.create_sheet("Lists")
ls["A1"], ls["B1"], ls["D1"] = "Stage", "SLA (days)", "Blocker reasons"
for k, (st, d) in enumerate(SLA.items()):
    ls[f"A{2 + k}"], ls[f"B{2 + k}"] = st, d
for k, b in enumerate(BLOCKERS[1:]):
    ls[f"D{2 + k}"] = b
ls.column_dimensions["A"].width = 20
ls.column_dimensions["D"].width = 40

db = wb.create_sheet("Dashboard", 0)
db["A1"] = "Weekly pipeline dashboard"
db["A1"].font = Font(name="Arial", size=14, bold=True, color=NAVY)
db["A2"] = '="As of "&TEXT(Tracker!$B$2,"dd-mmm-yy")'
db["A2"].font = F(italic=True, color="6B7280")
hdr = ["Stage", "Units in stage (WIP)", "Blocked", "Late vs. SLA", "Avg days in stage", "Rework loops (sum)"]
for j, h in enumerate(hdr, start=1):
    c = db.cell(row=4, column=j, value=h)
    c.font, c.fill = F(bold=True, color="FFFFFF"), HDR
    c.alignment = Alignment(wrap_text=True, horizontal="center")
R = f"Tracker!$E$5:$E${last}"
for k, st in enumerate(STAGES):
    r = 5 + k
    db[f"A{r}"] = st
    db[f"B{r}"] = f'=COUNTIF({R},A{r})'
    db[f"C{r}"] = f'=COUNTIFS({R},A{r},Tracker!$K$5:$K${last},"Blocked")'
    db[f"D{r}"] = f'=COUNTIFS({R},A{r},Tracker!$J$5:$J${last},">0")'
    db[f"E{r}"] = f'=IFERROR(AVERAGEIFS(Tracker!$I$5:$I${last},{R},A{r}),"")'
    db[f"F{r}"] = f'=SUMIFS(Tracker!$N$5:$N${last},{R},A{r})'
    db[f"E{r}"].number_format = "0.0"
for col, w in zip("ABCDEFGH", [34, 14, 12, 12, 12, 13, 14, 12]):
    db.column_dimensions[col].width = w
db.row_dimensions[4].height = 30

db["A13"] = "Weekly throughput vs. plan (target 30 finished videos / week)"
db["A13"].font = F(bold=True, color=NAVY)
for j, h in enumerate(["Week", "Week starting", "Target (cum.)", "Finished in week", "Finished (cum.)", "Gap vs. plan",
                       "Required / week to hit 300"], start=1):
    c = db.cell(row=14, column=j, value=h)
    c.font, c.fill = F(bold=True, color="FFFFFF"), HDR
    c.alignment = Alignment(wrap_text=True, horizontal="center")
db.row_dimensions[14].height = 30
for w in range(1, 11):
    r = 14 + w
    db[f"A{r}"] = w
    db[f"B{r}"] = W1 + dt.timedelta(weeks=w - 1)
    db[f"B{r}"].number_format = "dd-mmm"
    db[f"C{r}"] = 30 * w
    future = f"B{r}>Tracker!$B$2"
    db[f"D{r}"] = (f'=IF({future},"",COUNTIFS(Tracker!$O$5:$O${last},">="&B{r},'
                   f'Tracker!$O$5:$O${last},"<"&(B{r}+7)))')
    db[f"E{r}"] = f'=IF({future},"",SUM($D$15:D{r}))'
    db[f"F{r}"] = f'=IF({future},"",E{r}-C{r})'
    db[f"G{r}"] = f'=IF(OR({future},A{r}=10),"",ROUND((300-E{r})/(10-A{r}),1))'
db.conditional_formatting.add("F15:F24", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000", bold=True)))

db["A27"] = "Blocked units by reason and stage"
db["A27"].font = F(bold=True, color=NAVY)
for j, h in enumerate(["Blocker reason"] + STAGES[:5] + ["Total"], start=1):
    c = db.cell(row=28, column=j, value=h)
    c.font, c.fill = F(bold=True, color="FFFFFF"), HDR
    c.alignment = Alignment(wrap_text=True, horizontal="center")
db.row_dimensions[28].height = 30
for k, b in enumerate(BLOCKERS[1:]):
    r = 29 + k
    db[f"A{r}"] = f"=Lists!D{2 + k}"
    for j, st in enumerate(STAGES[:5]):
        col = "BCDEF"[j]
        db[f"{col}{r}"] = (f'=COUNTIFS({R},"{st}",Tracker!$L$5:$L${last},$A{r},'
                           f'Tracker!$K$5:$K${last},"Blocked")')
    db[f"G{r}"] = f"=SUM(B{r}:F{r})"

for sh in wb.worksheets:
    sh.sheet_view.showGridLines = sh.title != "Tracker"
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print("saved", OUT)
