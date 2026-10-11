"""Builds Luca_Step1_Financial_Model.xlsx with live formulas (Google Sheets compatible)."""
import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

import model

OUT = Path(__file__).resolve().parents[1] / "deliverables" / "Luca_Step1_Financial_Model.xlsx"

NAVY, TEAL, GREY, LIGHT = "1F2A44", "0F9D8A", "6B7280", "F3F4F6"
F_TITLE = Font(name="Arial", size=16, bold=True, color=NAVY)
F_SUB = Font(name="Arial", size=10, italic=True, color=GREY)
F_HDR = Font(name="Arial", size=10, bold=True, color="FFFFFF")
F_SEC = Font(name="Arial", size=10, bold=True, color=NAVY)
F_B = Font(name="Arial", size=10, bold=True)
F_N = Font(name="Arial", size=10)
F_IN = Font(name="Arial", size=10, color="0000FF")
F_LINK = Font(name="Arial", size=10, color="008000")
FILL_HDR = PatternFill("solid", fgColor=NAVY)
FILL_SEC = PatternFill("solid", fgColor="E5E7EB")
FILL_IN = PatternFill("solid", fgColor="FFF7D6")
FILL_TOT = PatternFill("solid", fgColor=LIGHT)
FILL_KPI = PatternFill("solid", fgColor="E6F6F3")
THIN = Side(style="thin", color="9CA3AF")
TOP = Border(top=THIN)
TOPBOT = Border(top=THIN, bottom=Side(style="double", color="9CA3AF"))

USD = '#,##0;(#,##0);"–"'
EUR = '€#,##0;(€#,##0);"–"'
PCT = '0.0%;(0.0%);"–"'
NUM1 = '0.0;(0.0);"–"'
INT = '0;(0);"–"'

M = model.MONTHS
MC = [L(3 + i) for i in range(M)]          # month columns C..N
FIRST, LAST, TOT = MC[0], MC[-1], L(3 + M)  # C, N, O


def style_range(ws, ref, font=None, fill=None, fmt=None, border=None, align=None):
    for row in ws[ref]:
        for c in row:
            if font: c.font = font
            if fill: c.fill = fill
            if fmt: c.number_format = fmt
            if border: c.border = border
            if align: c.alignment = align


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    ws["A2"] = sub
    ws["A2"].font = F_SUB
    ws.sheet_view.showGridLines = False


def month_header(ws, row=4, opening=True, total=True):
    """Row `row` = month index (used by formulas), row+1 = calendar month."""
    ws.cell(row=row, column=1, value="Month #").font = F_HDR
    ws.cell(row=row + 1, column=1, value="USD unless stated").font = F_HDR
    if opening:
        ws.cell(row=row, column=2, value="Opening").font = F_HDR
    for i, col in enumerate(MC):
        ws[f"{col}{row}"] = i + 1
        ws[f"{col}{row + 1}"] = f"=EDATE(Assumptions!$C$5,{col}${row}-1)"
        ws[f"{col}{row + 1}"].number_format = "mmm-yy"
    if total:
        ws[f"{TOT}{row}"] = "FY total"
    style_range(ws, f"A{row}:{TOT}{row + 1}", font=F_HDR, fill=FILL_HDR,
                align=Alignment(horizontal="center"))
    ws.cell(row=row, column=1).alignment = Alignment(horizontal="left")
    ws.cell(row=row + 1, column=1).alignment = Alignment(horizontal="left")
    ws.freeze_panes = ws[f"C{row + 2}"]
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 14
    for col in MC + [TOT]:
        ws.column_dimensions[col].width = 12.5


def section(ws, row, text):
    ws.cell(row=row, column=1, value=text)
    style_range(ws, f"A{row}:{TOT}{row}", font=F_SEC, fill=FILL_SEC)


def line(ws, row, label, fn, fmt=USD, total="sum", bold=False, border=None, opening=None, font=None):
    """fn(col, prev_col) -> formula for that month; prev_col is None for Month 1."""
    ws.cell(row=row, column=1, value=label).font = F_B if bold else F_N
    if opening is not None:
        ws[f"B{row}"] = opening
        ws[f"B{row}"].number_format = fmt
    for i, col in enumerate(MC):
        c = ws[f"{col}{row}"]
        c.value = fn(col, MC[i - 1] if i else None)
        c.number_format = fmt
        c.font = font or (F_B if bold else F_N)
    if total == "sum":
        ws[f"{TOT}{row}"] = f"=SUM({FIRST}{row}:{LAST}{row})"
    elif total == "last":
        ws[f"{TOT}{row}"] = f"={LAST}{row}"
    elif total:
        ws[f"{TOT}{row}"] = total
    ws[f"{TOT}{row}"].number_format = fmt
    ws[f"{TOT}{row}"].font = F_B
    ws[f"{TOT}{row}"].fill = FILL_TOT
    if border:
        style_range(ws, f"A{row}:{TOT}{row}", border=border)
    return row


# --------------------------------------------------------------------------------------------
wb = Workbook()
wb.calculation.fullCalcOnLoad = True
cover = wb.active
cover.title = "Cover"
A = wb.create_sheet("Assumptions")
H = wb.create_sheet("Headcount")
P = wb.create_sheet("P&L")
C = wb.create_sheet("Cash Flow")
R = wb.create_sheet("Runway")
LN = wb.create_sheet("Loan")
S = wb.create_sheet("Scenario Summary")

# ---------------------------------------------------------------- Assumptions
title(A, "Assumptions & scenario controls",
      "Blue cells on yellow = inputs. Everything else in the workbook is a formula. "
      "Source column shows whether a value comes from the case brief or is a Chief of Staff assumption.")
A.column_dimensions["A"].width = 3
A.column_dimensions["B"].width = 52
A.column_dimensions["C"].width = 16
A.column_dimensions["D"].width = 16
A.column_dimensions["E"].width = 16
A.column_dimensions["F"].width = 16
A.column_dimensions["G"].width = 70

ref = {}


def a_hdr(row, text, cols=("Value", "", "", "", "Source / rationale")):
    A.cell(row=row, column=2, value=text)
    for j, t in enumerate(cols):
        A.cell(row=row, column=3 + j, value=t)
    style_range(A, f"B{row}:G{row}", font=F_HDR, fill=FILL_HDR)


def a_in(row, key, label, value, fmt, src):
    A.cell(row=row, column=2, value=label).font = F_N
    c = A.cell(row=row, column=3, value=value)
    c.font, c.fill, c.number_format = F_IN, FILL_IN, fmt
    A.cell(row=row, column=7, value=src).font = F_SUB
    ref[key] = f"Assumptions!$C${row}"


a_hdr(4, "Model controls")
a_in(5, "start", "Model start month (Month 1)", dt.date(2026, 11, 1), "mmm-yy",
     "Assumption: first full month after the plan is approved.")
a_in(6, "sel", "Active operating scenario (1 = Base, 2 = Downside, 3 = Stress)", 1, "0",
     "Drives P&L, Cash Flow and Runway. Base = case assumptions exactly.")
a_in(7, "loan_on", "Show loan in P&L / Cash Flow statements? (1 = with loan, 0 = without)", 1, "0",
     "Runway tab always shows BOTH cases side by side, independent of this switch.")
A["D6"] = '=INDEX(D10:F10,1,C6)'
A["D6"].font = F_B
A["D7"] = '=IF(C7=1,"With €500k loan","Without loan")'
A["D7"].font = F_B
dv1 = DataValidation(type="whole", operator="between", formula1="1", formula2="3", allow_blank=False)
dv2 = DataValidation(type="whole", operator="between", formula1="0", formula2="1", allow_blank=False)
A.add_data_validation(dv1)
A.add_data_validation(dv2)
dv1.add("C6")
dv2.add("C7")

a_hdr(10, "Scenario table", ("Active", "Base", "Downside", "Stress", "Definition"))
A["B11"] = "Base = case brief exactly. Downside / Stress = Chief of Staff stress tests (illustrative)."
A["B11"].font = F_SUB
scen_rows = [
    ("g", "Revenue growth, MoM (compound)", "mom_growth", PCT,
     "Case: 6% MoM. Downside/Stress: growth stalls (flat at $1.5M)."),
    ("gm1", "Gross margin, Month 1", "gm_start", PCT, "Case: 65%."),
    ("gm12", "Gross margin, Month 12 (linear compression)", "gm_end", PCT,
     "Case: 60%. Downside/Stress: AI-video, hosting & content costs scale faster than revenue."),
    ("dso12", "Days sales outstanding (DSO), Month 12", "dso_end", "0",
     "Assumption: schools pay at 45 days today; Downside/Stress: collections slip to 90 days (public/"
     "private school payment cycles)."),
    ("extra", "Existing run-rate opex not in case brief (USD / month)", "extra_opex", USD,
     "Case brief contains only the 25 new hires + overhead + G&A. Stress layers an illustrative "
     "$1.0M/month for Luca's existing team, GTM, hosting & tools (~150 FTE)."),
]
r = 12
for key, label, attr, fmt, src in scen_rows:
    A.cell(row=r, column=2, value=label).font = F_N
    A[f"C{r}"] = f"=INDEX(D{r}:F{r},1,$C$6)"
    A[f"C{r}"].font, A[f"C{r}"].number_format = F_B, fmt
    fmt = "0.0%" if fmt == PCT else fmt
    A[f"C{r}"].number_format = fmt
    for j, sname in enumerate(["Base", "Downside", "Stress"]):
        c = A.cell(row=r, column=4 + j, value=getattr(model.SCENARIOS[sname], attr))
        c.font, c.fill, c.number_format = F_IN, FILL_IN, fmt
    A.cell(row=r, column=7, value=src).font = F_SUB
    ref[key] = f"Assumptions!$C${r}"
    r += 1

i = model.Inputs()
a_hdr(19, "Revenue")
a_in(20, "rev0", "Starting monthly revenue (Month 1)", i.start_revenue, USD, "Case: $1.5M.")

a_hdr(22, "Hiring & people costs")
a_in(23, "jr", "Junior salary (USD / month)", i.junior_salary, USD, "Case.")
a_in(24, "sr", "Senior salary (USD / month)", i.senior_salary, USD, "Case.")
a_in(25, "hire_cost", "One-off hiring cost per hire (USD)", i.hiring_cost, USD,
     "Case. Expensed in the hire's start month (recruiter fee / job boards / onboarding).")
a_in(26, "ben", "Benefits (% of total salary)", i.benefits_pct, PCT, "Case.")
a_in(27, "capex_hire", "Equipment capex per hire (USD, laptop + setup)", i.capex_per_hire, USD,
     "Assumption. Capitalised, depreciated straight-line.")
a_in(28, "dep_m", "Depreciation life (months)", i.dep_months, "0", "Assumption: 3 years.")
A["B29"] = "Hiring plan detail (25 hires, timing, level, function) lives on the Headcount tab."
A["B29"].font = F_SUB

a_hdr(31, "Operating expenses")
a_in(32, "ovh0", "Overhead, Month 1 (USD / month)", i.overhead_start, USD, "Case: $40k.")
a_in(33, "ovh_g", "Overhead growth over 12 months", i.overhead_growth_12m, PCT,
     "Case: +15% over 12 months. Applied LINEARLY: Month m = $40k × (1 + 15% × (m−1)/11) → "
     "$46.0k in Month 12 (+$545/month). Linear chosen because overhead steps with headcount/"
     "office additions, not compounding.")
a_in(34, "ga", "G&A / Legal (USD / month, fixed)", i.ga_legal, USD, "Case: $20k fixed.")

a_hdr(36, "Working capital, cash & tax")
a_in(37, "cash0", "Opening cash balance (USD)", i.opening_cash, USD,
     "Assumption (not in brief): $3.0M post-Series A cash on hand at Month 0.")
a_in(38, "dso0", "DSO at Month 0 / Month 1 (days)", i.dso_open, "0",
     "Assumption: schools invoiced monthly, paid ~45 days. AR = revenue × DSO / 30.")
a_in(39, "ap", "Payables days on COGS + overhead (days)", i.ap_days, "0",
     "Assumption: suppliers paid at 30 days. Opening AP set at Month-1 run-rate.")
a_in(40, "tax", "Corporate income tax rate (on positive EBT)", i.tax_rate, PCT,
     "Assumption: Mexico ISR 30%. Accrued monthly, paid the following month (provisional "
     "payments). No loss carry-forwards / PTU modelled.")

a_hdr(42, "Loan scenario")
a_in(43, "loan", "Loan amount (EUR)", i.loan_eur, EUR, "Case: €500k, drawn in full in Month 1.")
a_in(44, "rate", "Interest rate (annual, nominal)", i.loan_rate, PCT, "Case: 6% → 0.5% per month.")
a_in(45, "term", "Term (months)", i.loan_term, "0", "Case: 24 months.")
a_in(46, "io", "Interest-only period (months)", i.loan_io_months, "0",
     "Case: first 6 months interest-only, then annuity amortisation over remaining 18 months.")
a_in(47, "fx", "FX rate (USD per EUR)", i.fx, "0.00",
     "Case: €1 = $1.10, held flat. FX sensitivity shown on the Loan tab.")

S_ = ref  # shorthand

# ---------------------------------------------------------------- Headcount
title(H, "Headcount plan & hiring costs",
      "25 hires over Months 1–6 (10 senior / 15 junior). Edit start month or level in the table; "
      "everything downstream updates.")
hdr = ["#", "Role", "Function", "Level", "Start month", "Monthly salary", "Hiring cost",
       "Why this hire / what it unlocks"]
for j, t in enumerate(hdr):
    H.cell(row=4, column=1 + j, value=t)
style_range(H, "A4:O4", font=F_HDR, fill=FILL_HDR)
H.merge_cells("H4:O4")
why = {
    "Head of Content Operations": "Owns the 300-video pipeline, WIP limits and stage SLAs.",
    "Senior Account Executive (Mexico)": "Carries new-logo quota to sustain 6% MoM growth.",
    "Senior AI / Backend Engineer": "Builds the AI video-production tool (Step 3) to cut unit cost.",
    "Motion Designer": "Adds capacity at the slowest creative stage (5 h/unit).",
    "Scriptwriter": "Relieves the script bottleneck (4 h/unit, heavy rework).",
    "FP&A / Finance Manager": "Monthly close, cash forecasting, board reporting, AR collections.",
    "Senior Product Manager (AI Content Tools)": "Product owner for content tooling & AI roadmap.",
    "Customer Success Associate": "Keeps school onboarding + renewals ahead of customer growth.",
    "SDR / Business Development Rep": "Pipeline generation for AEs.",
    "Senior Account Executive (Expansion market)": "Opens first market outside Mexico.",
    "Senior Curriculum Lead (Secondary)": "Curriculum sign-off authority; reduces author↔script rework.",
    "Synthesia / AI Video Specialist": "Doubles final-assembly capacity (6 h/unit stage).",
    "Head of Customer Success": "Scales CS org; owns gross revenue retention.",
    "Graphic Designer": "Builds reusable asset library / templates.",
    "Data Analyst (RevOps / BI)": "Single source of truth for pipeline, revenue & production KPIs.",
    "Senior Full-stack Engineer": "Platform scalability for 2× users by Month 12.",
    "Implementation Specialist": "Faster school go-lives → faster revenue recognition & collection.",
    "Sound Editor": "Removes audio/video sync issues at final stage.",
    "Legal & Compliance Manager (International)": "Entity set-up, data privacy (minors), "
                                                  "contracts for international expansion.",
    "QA / Content Reviewer": "Pedagogical QA before release; catches errors before Stage 5.",
}
for k, (hid, role, func, lvl, start) in enumerate(model.HIRING_PLAN):
    rr = 5 + k
    H[f"A{rr}"] = hid
    H[f"B{rr}"] = role
    H[f"C{rr}"] = func
    H[f"D{rr}"] = lvl
    H[f"E{rr}"] = start
    for col in "DE":
        H[f"{col}{rr}"].font, H[f"{col}{rr}"].fill = F_IN, FILL_IN
    H[f"F{rr}"] = f'=IF(D{rr}="Senior",{S_["sr"]},{S_["jr"]})'
    H[f"G{rr}"] = f'={S_["hire_cost"]}'
    H[f"H{rr}"] = why[role]
    H[f"F{rr}"].number_format = H[f"G{rr}"].number_format = USD
    H[f"H{rr}"].font = F_SUB
    H.merge_cells(f"H{rr}:O{rr}")
HL = 5 + len(model.HIRING_PLAN) - 1  # last hire row (29)
H[f"B{HL + 1}"] = "Total"
H[f"B{HL + 1}"].font = F_B
H[f"E{HL + 1}"] = f"=COUNT(E5:E{HL})"
H[f"F{HL + 1}"] = f"=SUM(F5:F{HL})"
H[f"G{HL + 1}"] = f"=SUM(G5:G{HL})"
style_range(H, f"A{HL + 1}:H{HL + 1}", font=F_B, fill=FILL_TOT, border=TOP)
H[f"F{HL + 1}"].number_format = H[f"G{HL + 1}"].number_format = USD
dv3 = DataValidation(type="list", formula1='"Junior,Senior"')
dv4 = DataValidation(type="whole", operator="between", formula1="1", formula2="12")
H.add_data_validation(dv3)
H.add_data_validation(dv4)
dv3.add(f"D5:D{HL}")
dv4.add(f"E5:E{HL}")

g0 = HL + 4  # monthly grid header row
month_header(H, g0, opening=False)
H.freeze_panes = None
H.column_dimensions["A"].width = 6
H.column_dimensions["B"].width = 44
H.column_dimensions["C"].width = 22
H[f"A{g0}"] = ""
H[f"A{g0 + 1}"] = ""
H[f"B{g0}"] = "Month #"
H[f"B{g0 + 1}"] = "Headcount & people cost"
style_range(H, f"A{g0}:{TOT}{g0 + 1}", font=F_HDR, fill=FILL_HDR)
for col in MC:
    H[f"{col}{g0}"].alignment = Alignment(horizontal="center")
rng = lambda c: f"${c}$5:${c}${HL}"
rows_h = {}


def hline(row, label, fn, fmt=INT, total="sum", bold=False, border=None):
    H.cell(row=row, column=2, value=label).font = F_B if bold else F_N
    for col in MC:
        H[f"{col}{row}"] = fn(col)
        H[f"{col}{row}"].number_format = fmt
        H[f"{col}{row}"].font = F_B if bold else F_N
    if total == "sum":
        H[f"{TOT}{row}"] = f"=SUM({FIRST}{row}:{LAST}{row})"
    elif total == "last":
        H[f"{TOT}{row}"] = f"={LAST}{row}"
    H[f"{TOT}{row}"].number_format = fmt
    H[f"{TOT}{row}"].font = F_B
    H[f"{TOT}{row}"].fill = FILL_TOT
    if border:
        style_range(H, f"B{row}:{TOT}{row}", border=border)
    return row


r = g0 + 2
rows_h["new"] = hline(r, "New hires in month", lambda c: f"=COUNTIF({rng('E')},{c}${g0})")
rows_h["new_jr"] = hline(r + 1, "   of which junior",
                         lambda c: f'=COUNTIFS({rng("E")},{c}${g0},{rng("D")},"Junior")')
rows_h["new_sr"] = hline(r + 2, "   of which senior",
                         lambda c: f'=COUNTIFS({rng("E")},{c}${g0},{rng("D")},"Senior")')
rows_h["hc_jr"] = hline(r + 3, "Junior headcount (end of month)",
                        lambda c: f'=COUNTIFS({rng("E")},"<="&{c}${g0},{rng("D")},"Junior")', total="last")
rows_h["hc_sr"] = hline(r + 4, "Senior headcount (end of month)",
                        lambda c: f'=COUNTIFS({rng("E")},"<="&{c}${g0},{rng("D")},"Senior")', total="last")
rows_h["hc"] = hline(r + 5, "Total new-plan headcount", lambda c: f"={c}{r + 3}+{c}{r + 4}",
                     total="last", bold=True, border=TOP)
r2 = r + 7
H[f"B{r2}"] = "Headcount by function (end of month)"
H[f"B{r2}"].font = F_SEC
funcs = ["Content & Curriculum", "Sales & Partnerships", "Customer Success",
         "Product & Engineering", "Finance, Legal & Ops"]
for k, fname in enumerate(funcs):
    hline(r2 + 1 + k, f"   {fname}",
          lambda c, fname=fname: f'=COUNTIFS({rng("E")},"<="&{c}${g0},{rng("C")},"{fname}")',
          total="last")
r3 = r2 + 1 + len(funcs) + 1
H[f"B{r3}"] = "People cost (USD)"
H[f"B{r3}"].font = F_SEC
rows_h["sal"] = hline(r3 + 1, "Salaries", lambda c: f'=SUMIFS({rng("F")},{rng("E")},"<="&{c}${g0})', USD)
rows_h["ben"] = hline(r3 + 2, "Benefits (15% of salary)", lambda c: f"={c}{r3 + 1}*{S_['ben']}", USD)
rows_h["hire"] = hline(r3 + 3, "One-off hiring costs", lambda c: f"={c}{rows_h['new']}*{S_['hire_cost']}", USD)
rows_h["people"] = hline(r3 + 4, "Total people cost (P&L)", lambda c: f"=SUM({c}{r3 + 1}:{c}{r3 + 3})",
                         USD, bold=True, border=TOP)
rows_h["capex"] = hline(r3 + 5, "Equipment capex (cash, capitalised)",
                        lambda c: f"={c}{rows_h['new']}*{S_['capex_hire']}", USD)
rows_h["run"] = hline(r3 + 7, "Run-rate monthly cost per head (salary + benefits)",
                      lambda c: f"=IFERROR(({c}{r3 + 1}+{c}{r3 + 2})/{c}{rows_h['hc']},0)", USD, total=None)
hc = lambda key, col: f"Headcount!{col}{rows_h[key]}"

# ---------------------------------------------------------------- P&L
title(P, "Monthly Profit & Loss statement",
      "Scenario and loan treatment follow the controls on the Assumptions tab (shown in row 3).")
P["A3"] = '="Scenario: "&Assumptions!$D$6&"  |  "&Assumptions!$D$7'
P["A3"].font = Font(name="Arial", size=10, bold=True, color=TEAL)
month_header(P, 4, opening=False)
p = {}
r = 7
section(P, r, "Revenue & gross margin")
p["rev"] = line(P, r + 1, "Revenue", lambda c, pc: f"={S_['rev0']}*(1+{S_['g']})^({c}$4-1)", bold=True)
p["growth"] = line(P, r + 2, "   MoM growth",
                   lambda c, pc: f"=IF({c}$4=1,0,{c}{r + 1}/{pc}{r + 1}-1)", PCT, total=None)
p["gm"] = line(P, r + 3, "   Gross margin %",
               lambda c, pc: f"={S_['gm1']}+({S_['gm12']}-{S_['gm1']})*({c}$4-1)/11", PCT,
               total=f"={L(3 + M)}{r + 5}/{L(3 + M)}{r + 1}")
p["cogs"] = line(P, r + 4, "Cost of revenue (COGS)", lambda c, pc: f"=-{c}{r + 1}*(1-{c}{r + 3})")
p["gp"] = line(P, r + 5, "Gross profit", lambda c, pc: f"={c}{r + 1}+{c}{r + 4}", bold=True, border=TOP)
r = p["gp"] + 2
section(P, r, "Operating expenses")
p["sal"] = line(P, r + 1, "Salaries (new hires)", lambda c, pc: f"=-{hc('sal', c)}")
p["ben"] = line(P, r + 2, "Benefits (15% of salary)", lambda c, pc: f"=-{hc('ben', c)}")
p["hire"] = line(P, r + 3, "Hiring costs (one-off, $6k per hire)", lambda c, pc: f"=-{hc('hire', c)}")
p["ovh"] = line(P, r + 4, "Overhead (+15% linear over 12m)",
                lambda c, pc: f"=-{S_['ovh0']}*(1+{S_['ovh_g']}*({c}$4-1)/11)")
p["ga"] = line(P, r + 5, "G&A / Legal", lambda c, pc: f"=-{S_['ga']}")
p["extra"] = line(P, r + 6, "Existing run-rate opex (Stress scenario only)", lambda c, pc: f"=-{S_['extra']}")
p["opex"] = line(P, r + 7, "Total operating expenses", lambda c, pc: f"=SUM({c}{r + 1}:{c}{r + 6})",
                 bold=True, border=TOP)
p["opex_pct"] = line(P, r + 8, "   Opex as % of revenue", lambda c, pc: f"=-{c}{r + 7}/{c}{p['rev']}", PCT,
                     total=f"=-{TOT}{r + 7}/{TOT}{p['rev']}")
r = p["opex_pct"] + 2
p["ebitda"] = line(P, r, "EBITDA", lambda c, pc: f"={c}{p['gp']}+{c}{p['opex']}", bold=True, border=TOP)
p["ebitda_m"] = line(P, r + 1, "   EBITDA margin", lambda c, pc: f"={c}{r}/{c}{p['rev']}", PCT,
                     total=f"={TOT}{r}/{TOT}{p['rev']}")
p["dep"] = line(P, r + 2, "Depreciation (equipment)",
                lambda c, pc: f"=-SUM(Headcount!${FIRST}{rows_h['capex']}:Headcount!{c}{rows_h['capex']})/{S_['dep_m']}")
p["ebit"] = line(P, r + 3, "EBIT (operating profit)", lambda c, pc: f"={c}{r}+{c}{r + 2}", bold=True, border=TOP)
p["int"] = line(P, r + 4, "Interest expense – €500k loan",
                lambda c, pc: f"=-INDEX(Loan!$J$6:$J$29,{c}$4)*{S_['loan_on']}")
p["ebt"] = line(P, r + 5, "Profit before tax", lambda c, pc: f"={c}{r + 3}+{c}{r + 4}", bold=True)
p["tax"] = line(P, r + 6, "Income tax (30% on positive EBT)", lambda c, pc: f"=-MAX(0,{c}{r + 5})*{S_['tax']}")
p["ni"] = line(P, r + 7, "Net income", lambda c, pc: f"={c}{r + 5}+{c}{r + 6}", bold=True, border=TOPBOT)
p["ni_m"] = line(P, r + 8, "   Net margin", lambda c, pc: f"={c}{r + 7}/{c}{p['rev']}", PCT,
                 total=f"={TOT}{r + 7}/{TOT}{p['rev']}")
r = p["ni_m"] + 2
section(P, r, "Memo: both loan cases (feeds the Runway tab)")
p["int_wl"] = line(P, r + 1, "Interest expense if loan taken", lambda c, pc: f"=-INDEX(Loan!$J$6:$J$29,{c}$4)")
p["tax_nl"] = line(P, r + 2, "Income tax – without loan", lambda c, pc: f"=-MAX(0,{c}{p['ebit']})*{S_['tax']}")
p["tax_wl"] = line(P, r + 3, "Income tax – with loan",
                   lambda c, pc: f"=-MAX(0,{c}{p['ebit']}+{c}{r + 1})*{S_['tax']}")
p["ni_nl"] = line(P, r + 4, "Net income – without loan", lambda c, pc: f"={c}{p['ebit']}+{c}{r + 2}", bold=True)
p["ni_wl"] = line(P, r + 5, "Net income – with loan",
                  lambda c, pc: f"={c}{p['ebit']}+{c}{r + 1}+{c}{r + 3}", bold=True)
p["ni_diff"] = line(P, r + 6, "After-tax cost of the loan (P&L)", lambda c, pc: f"={c}{r + 5}-{c}{r + 4}")
r = p["ni_diff"] + 2
section(P, r, "Cost driver view (share of total costs before tax)")
cost_rows = [("COGS", "cogs"), ("People (salary + benefits + hiring)", None), ("Overhead", "ovh"),
             ("G&A / Legal", "ga"), ("Existing run-rate opex", "extra")]
tot_cost = (lambda c: f"(-{c}{p['cogs']}-{c}{p['opex']}-{c}{p['dep']})")
for k, (lab, key) in enumerate(cost_rows):
    if key:
        fn = (lambda c, pc, key=key: f"=-{c}{p[key]}/{tot_cost(c)}")
        tfn = f"=-{TOT}{p[key]}/{tot_cost(TOT)}"
    else:
        fn = (lambda c, pc: f"=-({c}{p['sal']}+{c}{p['ben']}+{c}{p['hire']})/{tot_cost(c)}")
        tfn = f"=-({TOT}{p['sal']}+{TOT}{p['ben']}+{TOT}{p['hire']})/{tot_cost(TOT)}"
    line(P, r + 1 + k, f"   {lab}", fn, PCT, total=tfn)
P.conditional_formatting.add(f"C{p['ni']}:{TOT}{p['ni']}",
                             CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000", bold=True)))

# ---------------------------------------------------------------- Cash Flow (indirect)
title(C, "Cash flow statement – indirect method",
      "Starts from net income, adds back non-cash items and working-capital movements. "
      "Loan treatment follows Assumptions!C7.")
C["A3"] = '="Scenario: "&Assumptions!$D$6&"  |  "&Assumptions!$D$7'
C["A3"].font = Font(name="Arial", size=10, bold=True, color=TEAL)
month_header(C, 4)
cf = {}
# supporting balances first (lower on the sheet) so references are known
sb = 32
section(C, sb, "Supporting schedule – working capital balances (end of month)")
cf["dso"] = line(C, sb + 1, "DSO (days, linear path to scenario Month-12 value)",
                 lambda c, pc: f"={S_['dso0']}+({S_['dso12']}-{S_['dso0']})*({c}$4-1)/11", NUM1,
                 total=None, opening=f"={S_['dso0']}")
cf["ar"] = line(C, sb + 2, "Accounts receivable (revenue × DSO / 30)",
                lambda c, pc: f"='P&L'!{c}{p['rev']}*{c}{sb + 1}/30", total="last",
                opening=f"={S_['rev0']}*{S_['dso0']}/30")
cf["ap"] = line(C, sb + 3, "Accounts payable ((COGS + overhead) × AP days / 30)",
                lambda c, pc: f"=-('P&L'!{c}{p['cogs']}+'P&L'!{c}{p['ovh']})*{S_['ap']}/30", total="last",
                opening=f"=-('P&L'!{FIRST}{p['cogs']}+'P&L'!{FIRST}{p['ovh']})*{S_['ap']}/30")
cf["taxp"] = line(C, sb + 4, "Income tax payable (current month, paid next month)",
                  lambda c, pc: f"=-'P&L'!{c}{p['tax']}", total="last",
                  opening=f"=-'P&L'!{FIRST}{p['tax']}")
cf["ppe"] = line(C, sb + 5, "Net PP&E (cumulative capex − depreciation)",
                 lambda c, pc: f"={'B' if pc is None else pc}{sb + 5}+Headcount!{c}{rows_h['capex']}+'P&L'!{c}{p['dep']}",
                 total="last", opening=0)
cf["loanbal"] = line(C, sb + 6, "Loan balance (USD)",
                     lambda c, pc: f"=INDEX(Loan!$L$6:$L$29,{c}$4)*{S_['loan_on']}", total="last", opening=0)

r = 7
section(C, r, "Operating activities")
cf["ni"] = line(C, r + 1, "Net income", lambda c, pc: f"='P&L'!{c}{p['ni']}", bold=True)
cf["dep"] = line(C, r + 2, "+ Depreciation (non-cash)", lambda c, pc: f"=-'P&L'!{c}{p['dep']}")
prev = lambda c, pc, row: f"{'B' if pc is None else pc}{row}"
cf["dar"] = line(C, r + 3, "− Increase in accounts receivable",
                 lambda c, pc: f"=-({c}{cf['ar']}-{prev(c, pc, cf['ar'])})")
cf["dap"] = line(C, r + 4, "+ Increase in accounts payable",
                 lambda c, pc: f"={c}{cf['ap']}-{prev(c, pc, cf['ap'])}")
cf["dtax"] = line(C, r + 5, "+ Increase in income tax payable",
                  lambda c, pc: f"={c}{cf['taxp']}-{prev(c, pc, cf['taxp'])}")
cf["cfo"] = line(C, r + 6, "Cash flow from operations (CFO)", lambda c, pc: f"=SUM({c}{r + 1}:{c}{r + 5})",
                 bold=True, border=TOP)
r = cf["cfo"] + 2
section(C, r, "Investing activities")
cf["capex"] = line(C, r + 1, "Equipment capex for new hires", lambda c, pc: f"=-Headcount!{c}{rows_h['capex']}")
cf["cfi"] = line(C, r + 2, "Cash flow from investing (CFI)", lambda c, pc: f"={c}{r + 1}", bold=True, border=TOP)
r = cf["cfi"] + 2
section(C, r, "Financing activities")
cf["draw"] = line(C, r + 1, "Loan drawdown (€500k × 1.10)",
                  lambda c, pc: f"=IF({c}$4=1,{S_['loan']}*{S_['fx']},0)*{S_['loan_on']}")
cf["princ"] = line(C, r + 2, "Loan principal repayment", lambda c, pc: f"=-INDEX(Loan!$I$6:$I$29,{c}$4)*{S_['loan_on']}")
cf["cff"] = line(C, r + 3, "Cash flow from financing (CFF)", lambda c, pc: f"={c}{r + 1}+{c}{r + 2}",
                 bold=True, border=TOP)
r = cf["cff"] + 2
cf["net"] = line(C, r, "Net change in cash", lambda c, pc: f"={c}{cf['cfo']}+{c}{cf['cfi']}+{c}{cf['cff']}",
                 bold=True, border=TOP)
cf["open"] = line(C, r + 1, "Opening cash",
                  lambda c, pc: f"={S_['cash0']}" if pc is None else f"={pc}{r + 2}", total=f"={FIRST}{r + 1}")
cf["close"] = line(C, r + 2, "Closing cash", lambda c, pc: f"={c}{r + 1}+{c}{r}", bold=True, border=TOPBOT,
                   total="last")
cf["fcf"] = line(C, r + 4, "Memo: free cash flow (CFO + CFI, pre-financing)",
                 lambda c, pc: f"={c}{cf['cfo']}+{c}{cf['cfi']}")
cf["burn"] = line(C, r + 5, "Memo: net burn (positive = burning cash)", lambda c, pc: f"=MAX(0,-{c}{cf['fcf']})")
assert cf["burn"] < sb, "cash flow rows overlap supporting schedule"

bs = sb + 8
section(C, bs, "Balance sheet integrity check (assets − liabilities − equity must be zero)")
cf["assets"] = line(C, bs + 1, "Total assets (cash + AR + net PP&E)",
                    lambda c, pc: f"={c}{cf['close']}+{c}{cf['ar']}+{c}{cf['ppe']}", total="last",
                    opening=f"={S_['cash0']}+B{cf['ar']}+B{cf['ppe']}")
cf["liab"] = line(C, bs + 2, "Total liabilities (AP + tax payable + loan)",
                  lambda c, pc: f"={c}{cf['ap']}+{c}{cf['taxp']}+{c}{cf['loanbal']}", total="last",
                  opening=f"=B{cf['ap']}+B{cf['taxp']}+B{cf['loanbal']}")
cf["eq"] = line(C, bs + 3, "Equity (opening + cumulative net income)",
                lambda c, pc: f"={prev(c, pc, bs + 3)}+{c}{cf['ni']}", total="last",
                opening=f"=B{bs + 1}-B{bs + 2}")
cf["chk"] = line(C, bs + 4, "Check (should be 0)",
                 lambda c, pc: f"=ROUND({c}{bs + 1}-{c}{bs + 2}-{c}{bs + 3},2)", total="last", bold=True)
C.conditional_formatting.add(f"C{bs + 4}:{TOT}{bs + 4}",
                             CellIsRule(operator="notEqual", formula=["0"], fill=PatternFill("solid", fgColor="FFC7CE")))
C.conditional_formatting.add(f"C{cf['close']}:{TOT}{cf['close']}",
                             CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000", bold=True)))

# ---------------------------------------------------------------- Loan
title(LN, "€500k loan – 24-month repayment waterfall",
      "6% p.a. (0.5%/month). Months 1–6 interest-only; Months 7–24 equal annuity payments. "
      "EUR schedule converted at the Assumptions FX rate.")
heads = ["Month", "Date", "Phase", "Opening balance (€)", "Interest (€)", "Principal (€)",
         "Total payment (€)", "Closing balance (€)", "Principal ($)", "Interest ($)", "Total payment ($)",
         "Closing balance ($)"]
for j, t in enumerate(heads):
    LN.cell(row=5, column=1 + j, value=t)
style_range(LN, "A5:L5", font=F_HDR, fill=FILL_HDR, align=Alignment(horizontal="center", wrap_text=True))
LN.row_dimensions[5].height = 30
for j, w in enumerate([8, 10, 16, 18, 14, 14, 16, 18, 14, 14, 16, 18]):
    LN.column_dimensions[L(1 + j)].width = w
pmt = f"-PMT({S_['rate']}/12,{S_['term']}-{S_['io']},{S_['loan']})"
for m in range(1, 25):
    rr = 5 + m
    LN[f"A{rr}"] = m
    LN[f"B{rr}"] = f"=EDATE({S_['start']},A{rr}-1)"
    LN[f"B{rr}"].number_format = "mmm-yy"
    LN[f"C{rr}"] = f'=IF(A{rr}<={S_["io"]},"Interest-only","Amortising")'
    LN[f"D{rr}"] = f"={S_['loan']}" if m == 1 else f"=H{rr - 1}"
    LN[f"E{rr}"] = f"=D{rr}*{S_['rate']}/12"
    LN[f"G{rr}"] = f"=IF(A{rr}<={S_['io']},E{rr},{pmt})"
    LN[f"F{rr}"] = f"=G{rr}-E{rr}"
    LN[f"H{rr}"] = f"=D{rr}-F{rr}"
    LN[f"I{rr}"] = f"=F{rr}*{S_['fx']}"
    LN[f"J{rr}"] = f"=E{rr}*{S_['fx']}"
    LN[f"K{rr}"] = f"=G{rr}*{S_['fx']}"
    LN[f"L{rr}"] = f"=H{rr}*{S_['fx']}"
    for col in "DEFGH":
        LN[f"{col}{rr}"].number_format = EUR
    for col in "IJKL":
        LN[f"{col}{rr}"].number_format = USD
    if m == 12:
        style_range(LN, f"A{rr}:L{rr}", border=Border(bottom=Side(style="medium", color=TEAL)))
LN["C30"] = "Total (24 months)"
for col in "EFGIJK":
    LN[f"{col}30"] = f"=SUM({col}6:{col}29)"
    LN[f"{col}30"].number_format = EUR if col in "EFG" else USD
LN["C31"] = "Of which inside 12-month model"
for col in "EFGIJK":
    LN[f"{col}31"] = f"=SUM({col}6:{col}17)"
    LN[f"{col}31"].number_format = EUR if col in "EFG" else USD
style_range(LN, "A30:L31", font=F_B, fill=FILL_TOT)
LN["A33"] = "Key loan facts"
LN["A33"].font = F_SEC
facts = [
    ("Monthly payment, Months 1–6 (interest-only, €)", "=G6", EUR),
    ("Monthly annuity payment, Months 7–24 (€)", "=G12", EUR),
    ("Monthly annuity payment, Months 7–24 ($)", "=K12", USD),
    ("Total interest over 24 months ($)", "=J30", USD),
    ("Balance still outstanding at Month 12 ($)", "=L17", USD),
    ("Cash received at drawdown ($)", f"={S_['loan']}*{S_['fx']}", USD),
]
for k, (lab, f, fmt) in enumerate(facts):
    LN[f"A{34 + k}"] = lab
    LN[f"H{34 + k}"] = f
    LN[f"H{34 + k}"].number_format = fmt
    LN[f"H{34 + k}"].font = F_B
LN["A41"] = "FX sensitivity – USD cost of servicing the EUR loan (revenue is in MXN/USD)"
LN["A41"].font = F_SEC
LN["A42"], LN["D42"], LN["G42"] = "USD per EUR", "Total repaid, 24m ($)", "Extra cost vs. base ($)"
style_range(LN, "A42:L42", font=F_HDR, fill=FILL_HDR)
for k, fxv in enumerate([1.00, 1.10, 1.20, 1.30]):
    rr = 43 + k
    LN[f"A{rr}"] = fxv
    LN[f"A{rr}"].number_format = "0.00"
    LN[f"A{rr}"].font, LN[f"A{rr}"].fill = F_IN, FILL_IN
    LN[f"D{rr}"] = f"=$G$30*A{rr}"
    LN[f"G{rr}"] = f"=D{rr}-$G$30*{S_['fx']}"
    LN[f"D{rr}"].number_format = LN[f"G{rr}"].number_format = USD
LN.freeze_panes = "A6"
LN.sheet_view.showGridLines = False

# ---------------------------------------------------------------- Runway
title(R, "Cash burn & runway – with vs. without the €500k loan",
      "Both cases computed in parallel for the active operating scenario (Assumptions!C6). "
      "Runway = months until cash < 0; beyond Month 12 extrapolated on the average burn of Months 10–12.")
R["A3"] = '="Operating scenario: "&Assumptions!$D$6'
R["A3"].font = Font(name="Arial", size=10, bold=True, color=TEAL)
month_header(R, 4)
rw = {}


def cash_block(start, label, ni_row, tax_row, with_loan):
    k = {}
    section(R, start, label)
    k["ni"] = line(R, start + 1, "Net income", lambda c, pc: f"='P&L'!{c}{ni_row}")
    k["wc"] = line(R, start + 2, "+ D&A and working-capital movements (AR, AP)",
                   lambda c, pc: f"='Cash Flow'!{c}{cf['dep']}+'Cash Flow'!{c}{cf['dar']}+'Cash Flow'!{c}{cf['dap']}")
    k["dtax"] = line(R, start + 3, "+ Increase in tax payable",
                     lambda c, pc: f"=-'P&L'!{c}{tax_row}+'P&L'!{pc}{tax_row}" if pc else "=0")
    k["capex"] = line(R, start + 4, "− Capex", lambda c, pc: f"='Cash Flow'!{c}{cf['capex']}")
    k["fcf"] = line(R, start + 5, "Free cash flow (pre-financing)", lambda c, pc: f"=SUM({c}{start + 1}:{c}{start + 4})",
                    bold=True, border=TOP)
    if with_loan:
        k["loan"] = line(R, start + 6, "+ Loan drawdown − principal repayment",
                         lambda c, pc: f"=IF({c}$4=1,{S_['loan']}*{S_['fx']},0)-INDEX(Loan!$I$6:$I$29,{c}$4)")
    else:
        k["loan"] = line(R, start + 6, "Loan cash flows", lambda c, pc: "=0")
    k["net"] = line(R, start + 7, "Net cash flow", lambda c, pc: f"={c}{start + 5}+{c}{start + 6}", bold=True)
    k["burn"] = line(R, start + 8, "Monthly net burn (positive = burning; negative = generating)",
                     lambda c, pc: f"=-{c}{start + 7}")
    k["cash"] = line(R, start + 9, "Closing cash", lambda c, pc: f"={S_['cash0']}+{c}{start + 7}" if pc is None
                     else f"={pc}{start + 9}+{c}{start + 7}", bold=True, border=TOPBOT, total="last",
                     opening=f"={S_['cash0']}")
    k["cum"] = line(R, start + 10, "Cumulative net cash flow vs. Month 0",
                    lambda c, pc: f"={c}{start + 9}-{S_['cash0']}", total="last")
    k["debt"] = line(R, start + 11, "Loan balance outstanding",
                     (lambda c, pc: f"=INDEX(Loan!$L$6:$L$29,{c}$4)") if with_loan else (lambda c, pc: "=0"),
                     total="last")
    k["netcash"] = line(R, start + 12, "Net cash (cash − debt)", lambda c, pc: f"={c}{start + 9}-{c}{start + 11}",
                        total="last")
    k["runout"] = line(R, start + 13, "Helper: cash-out point (months, interpolated)",
                       lambda c, pc: (f"=IF(AND({'B' if pc is None else pc}{start + 9}>=0,{c}{start + 9}<0),"
                                      f"({c}$4-1)+{'B' if pc is None else pc}{start + 9}/"
                                      f"({'B' if pc is None else pc}{start + 9}-{c}{start + 9}),\"\")"),
                       NUM1, total=None, font=F_SUB)
    R.conditional_formatting.add(f"C{k['cash']}:{TOT}{k['cash']}",
                                 CellIsRule(operator="lessThan", formula=["0"],
                                            font=Font(color="C00000", bold=True),
                                            fill=PatternFill("solid", fgColor="FFC7CE")))
    return k


rw["nl"] = cash_block(7, "Case A – Base plan WITHOUT loan", p["ni_nl"], p["tax_nl"], False)
rw["wl"] = cash_block(23, "Case B – Base plan WITH €500k loan", p["ni_wl"], p["tax_wl"], True)

kp = 40
for j, t in enumerate(["Runway & burn KPIs", "", "Without loan", "With loan", "Loan impact"]):
    R.cell(row=kp, column=1 + j, value=t)
style_range(R, f"A{kp}:E{kp}", font=F_HDR, fill=FILL_HDR)


def kpi(row, label, f_nl, f_wl, fmt, diff=True):
    R[f"A{row}"] = label
    R[f"C{row}"] = f_nl
    R[f"D{row}"] = f_wl
    if diff:
        R[f"E{row}"] = f'=IFERROR(D{row}-C{row},"n/a")'
    for col in "CDE":
        R[f"{col}{row}"].number_format = fmt
        R[f"{col}{row}"].font = F_B
        R[f"{col}{row}"].fill = FILL_KPI


rng_ = lambda k, key: f"${FIRST}${rw[k][key]}:${LAST}${rw[k][key]}"
avg_burn = lambda k: f"AVERAGE(${L(3 + M - 3)}${rw[k]['burn']}:${LAST}${rw[k]['burn']})"


def runway_f(k):
    return (f'=IF(COUNT({rng_(k, "runout")})>0,MIN({rng_(k, "runout")}),'
            f'IF({avg_burn(k)}>0,12+${LAST}${rw[k]["cash"]}/{avg_burn(k)},"Self-funding"))')


kpi(kp + 1, "Opening cash (Month 0)", f"={S_['cash0']}", f"={S_['cash0']}", USD)
kpi(kp + 2, "Cash at Month 12", f"=${LAST}${rw['nl']['cash']}", f"=${LAST}${rw['wl']['cash']}", USD)
kpi(kp + 3, "Lowest month-end cash in 12 months", f"=MIN({rng_('nl', 'cash')})", f"=MIN({rng_('wl', 'cash')})", USD)
kpi(kp + 4, "Month of lowest cash",
    f"=MATCH(C{kp + 3},{rng_('nl', 'cash')},0)", f"=MATCH(D{kp + 3},{rng_('wl', 'cash')},0)", "0", diff=False)
kpi(kp + 5, "Months with negative free cash flow (of 12)",
    f'=COUNTIF({rng_("nl", "fcf")},"<0")', f'=COUNTIF({rng_("wl", "fcf")},"<0")', "0")
kpi(kp + 6, "Average monthly net burn, Months 10–12 (negative = generating)",
    f"={avg_burn('nl')}", f"={avg_burn('wl')}", USD)
kpi(kp + 7, "Cumulative 12-month net burn (negative = cash generated)",
    f"=-${LAST}${rw['nl']['cum']}", f"=-${LAST}${rw['wl']['cum']}", USD)
kpi(kp + 8, "RUNWAY (months until cash < 0)", runway_f("nl"), runway_f("wl"), NUM1)
kpi(kp + 9, "Net cash at Month 12 (cash − debt)", f"=${LAST}${rw['nl']['netcash']}",
    f"=${LAST}${rw['wl']['netcash']}", USD)
kpi(kp + 10, "Total loan interest paid within 12 months", "=0", f"=Loan!J31", USD)
R[f"A{kp + 8}"].font = F_B
R[f"A{kp + 12}"] = ("Reading guide: if runway shows \"Self-funding\" the plan never burns cash in Months 10–12 "
                    "and cash never goes negative, i.e. runway is not a binding constraint. Switch Assumptions!C6 "
                    "to 3 (Stress) to see a burning scenario.")
R[f"A{kp + 12}"].font = F_SUB
R[f"A{kp + 14}"] = "Integrity check: Runway tab cash (selected loan case) − Cash Flow tab closing cash"
R[f"C{kp + 14}"] = (f"=ROUND(IF({S_['loan_on']}=1,${LAST}${rw['wl']['cash']},${LAST}${rw['nl']['cash']})"
                    f"-'Cash Flow'!${LAST}${cf['close']},2)")
R[f"C{kp + 14}"].number_format = USD
R[f"C{kp + 14}"].font = F_B
R.column_dimensions["A"].width = 58

ch = LineChart()
ch.title = "Month-end cash: with vs. without loan (USD)"
ch.height, ch.width = 8, 22
ch.y_axis.title = "USD"
ch.y_axis.numFmt = '#,##0'
for k in ("nl", "wl"):
    ch.add_data(Reference(R, min_col=3, max_col=3 + M - 1, min_row=rw[k]["cash"]), from_rows=True,
                titles_from_data=False)
ch.series[0].tx = SeriesLabel(v="Without loan")
ch.series[1].tx = SeriesLabel(v="With €500k loan")
ch.set_categories(Reference(R, min_col=3, max_col=3 + M - 1, min_row=5))
R.add_chart(ch, f"G{kp}")

# ---------------------------------------------------------------- Cover / dashboard
title(cover, "Luca – 12-month financial model (Chief of Staff hiring case)",
      "All data fictional. Model is fully formula-driven; change inputs on the Assumptions tab.")
cover.column_dimensions["A"].width = 3
cover.column_dimensions["B"].width = 58
for col in "CDEF":
    cover.column_dimensions[col].width = 18
cover["B4"] = "How to use"
cover["B4"].font = F_SEC
guide = [
    "1. Assumptions – every input (blue on yellow), scenario selector (C6) and loan switch (C7).",
    "2. Headcount – 25 named hires (role, level, start month) → monthly headcount, salaries, benefits, hiring cost.",
    "3. P&L – monthly profit & loss; memo block shows net income with AND without the loan.",
    "4. Cash Flow – indirect method (net income → CFO → CFI → CFF), working-capital schedule, balance-sheet check.",
    "5. Runway – monthly & cumulative cash, burn and runway for Base WITH loan vs. WITHOUT loan, side by side.",
    "6. Loan – €500k, 6%, 24m (6m interest-only + 18m annuity) waterfall in EUR and USD + FX sensitivity.",
    "7. Scenario Summary – Base / Downside / Stress outputs side by side (snapshot) for the board deck.",
]
for k, g in enumerate(guide):
    cover[f"B{5 + k}"] = g
    cover[f"B{5 + k}"].font = F_N
cover["B13"] = "Colour code: blue font on yellow = input · black = formula · green = link to another tab."
cover["B13"].font = F_SUB

cover["B15"] = '="Headline KPIs – active scenario: "&Assumptions!$D$6'
cover["B15"].font = F_SEC
for j, t in enumerate(["Metric", "Month 1", "Month 12", "12-month total", ""]):
    cover.cell(row=16, column=2 + j, value=t)
style_range(cover, "B16:E16", font=F_HDR, fill=FILL_HDR)
dash = [
    ("Revenue", f"='P&L'!{FIRST}{p['rev']}", f"='P&L'!{LAST}{p['rev']}", f"='P&L'!{TOT}{p['rev']}", USD),
    ("Gross margin %", f"='P&L'!{FIRST}{p['gm']}", f"='P&L'!{LAST}{p['gm']}", f"='P&L'!{TOT}{p['gm']}", PCT),
    ("Total operating expenses", f"=-'P&L'!{FIRST}{p['opex']}", f"=-'P&L'!{LAST}{p['opex']}",
     f"=-'P&L'!{TOT}{p['opex']}", USD),
    ("EBITDA", f"='P&L'!{FIRST}{p['ebitda']}", f"='P&L'!{LAST}{p['ebitda']}", f"='P&L'!{TOT}{p['ebitda']}", USD),
    ("EBITDA margin", f"='P&L'!{FIRST}{p['ebitda_m']}", f"='P&L'!{LAST}{p['ebitda_m']}",
     f"='P&L'!{TOT}{p['ebitda_m']}", PCT),
    ("Net income (statement case)", f"='P&L'!{FIRST}{p['ni']}", f"='P&L'!{LAST}{p['ni']}", f"='P&L'!{TOT}{p['ni']}", USD),
    ("New-plan headcount (end of month)", f"=Headcount!{FIRST}{rows_h['hc']}", f"=Headcount!{LAST}{rows_h['hc']}",
     f"=Headcount!{TOT}{rows_h['hc']}", INT),
    ("Free cash flow (pre-financing)", f"='Cash Flow'!{FIRST}{cf['fcf']}", f"='Cash Flow'!{LAST}{cf['fcf']}",
     f"='Cash Flow'!{TOT}{cf['fcf']}", USD),
    ("Closing cash – WITHOUT loan", f"=Runway!{FIRST}{rw['nl']['cash']}", f"=Runway!{LAST}{rw['nl']['cash']}", "", USD),
    ("Closing cash – WITH loan", f"=Runway!{FIRST}{rw['wl']['cash']}", f"=Runway!{LAST}{rw['wl']['cash']}", "", USD),
    ("Runway – WITHOUT loan (months)", "", "", f"=Runway!C{kp + 8}", NUM1),
    ("Runway – WITH loan (months)", "", "", f"=Runway!D{kp + 8}", NUM1),
    ("Balance-sheet check (must be 0)", "", "", f"='Cash Flow'!{TOT}{cf['chk']}", USD),
]
for k, (lab, f1, f12, ft, fmt) in enumerate(dash):
    rr = 17 + k
    cover[f"B{rr}"] = lab
    for col, f in zip("CDE", (f1, f12, ft)):
        cover[f"{col}{rr}"] = f if f else None
        cover[f"{col}{rr}"].number_format = fmt
        cover[f"{col}{rr}"].font = F_LINK
    if k % 2 == 0:
        style_range(cover, f"B{rr}:E{rr}", fill=FILL_TOT)

bc = BarChart()
bc.title = "Revenue and EBITDA (USD / month)"
bc.height, bc.width = 8, 18
bc.add_data(Reference(P, min_col=3, max_col=3 + M - 1, min_row=p["rev"]), from_rows=True)
bc.add_data(Reference(P, min_col=3, max_col=3 + M - 1, min_row=p["ebitda"]), from_rows=True)
bc.series[0].tx = SeriesLabel(v="Revenue")
bc.series[1].tx = SeriesLabel(v="EBITDA")
bc.set_categories(Reference(P, min_col=3, max_col=3 + M - 1, min_row=5))
bc.y_axis.numFmt = '#,##0'
cover.add_chart(bc, "G4")

# ---------------------------------------------------------------- Scenario summary (snapshot)
title(S, "Scenario summary – Base vs. Downside vs. Stress (snapshot values)",
      "Static values generated from the same model logic for all three scenarios at once (Sheets has no "
      "data tables). To see any scenario live, set Assumptions!C6.")
S.column_dimensions["A"].width = 3
S.column_dimensions["B"].width = 78
for col in "CDEFGH":
    S.column_dimensions[col].width = 17
cols = ["Base", "Downside", "Stress"]
for j, t in enumerate(["Metric"] + cols):
    S.cell(row=4, column=2 + j, value=t)
style_range(S, "B4:E4", font=F_HDR, fill=FILL_HDR)
res = {n: model.run(model.SCENARIOS[n]) for n in cols}


def rw_or(x):
    return "Self-funding" if x is None else round(x, 1)


metrics = [
    ("MoM revenue growth", lambda n, o: model.SCENARIOS[n].mom_growth, PCT),
    ("Gross margin, Month 12", lambda n, o: o["gm"][-1], PCT),
    ("DSO, Month 12 (days)", lambda n, o: o["dso"][-1], "0"),
    ("Existing run-rate opex added (USD / month)", lambda n, o: o["extra"][-1], USD),
    ("Revenue – 12-month total", lambda n, o: sum(o["revenue"]), USD),
    ("Revenue – Month 12", lambda n, o: o["revenue"][-1], USD),
    ("EBITDA – 12-month total", lambda n, o: sum(o["ebitda"]), USD),
    ("Net income – 12m, without loan", lambda n, o: sum(o["ni_nl"]), USD),
    ("Net income – 12m, with loan", lambda n, o: sum(o["ni_wl"]), USD),
    ("Cash at Month 12 – without loan", lambda n, o: o["cash_nl"][-1], USD),
    ("Cash at Month 12 – with loan", lambda n, o: o["cash_wl"][-1], USD),
    ("Lowest month-end cash – without loan", lambda n, o: min(o["cash_nl"]), USD),
    ("Lowest month-end cash – with loan", lambda n, o: min(o["cash_wl"]), USD),
    ("Runway (months) – without loan", lambda n, o: rw_or(model.runway(o["cash_nl"], o["net_cf_nl"])), NUM1),
    ("Runway (months) – with loan", lambda n, o: rw_or(model.runway(o["cash_wl"], o["net_cf_wl"])), NUM1),
    ("Extra fixed cost absorbable before cash < 0 in 12m – without loan (USD/month)",
     lambda n, o: model.headroom(model.SCENARIOS[n]), USD),
    ("Extra fixed cost absorbable before cash < 0 in 12m – with loan (USD/month)",
     lambda n, o: model.headroom(model.SCENARIOS[n], "cash_wl"), USD),
]
for k, (lab, fn, fmt) in enumerate(metrics):
    rr = 5 + k
    S[f"B{rr}"] = lab
    for j, n in enumerate(cols):
        c = S.cell(row=rr, column=3 + j, value=fn(n, res[n]))
        c.number_format = fmt
        c.alignment = Alignment(horizontal="right")
    if k % 2 == 0:
        style_range(S, f"B{rr}:E{rr}", fill=FILL_TOT)
S[f"B{6 + len(metrics)}"] = ("Stress = Downside (flat revenue, GM → 55%, DSO → 90 days) PLUS an illustrative $1.0M/month "
                             "existing cost base that the case brief does not include. It is a reverse stress test, "
                             "not a forecast.")
S[f"B{6 + len(metrics)}"].font = F_SUB

for ws in wb.worksheets:
    ws.sheet_properties.tabColor = {"Cover": TEAL, "Assumptions": "F59E0B"}.get(ws.title, NAVY)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name != "Arial":
                c.font = Font(name="Arial", size=c.font.size if c.font else 10, bold=c.font.bold,
                              italic=c.font.italic, color=c.font.color)

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print("saved", OUT)
print("rows", dict(p=p, cf=cf, rw={k: v["cash"] for k, v in rw.items()}, kp=kp))
