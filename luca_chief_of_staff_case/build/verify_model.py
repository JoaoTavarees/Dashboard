"""Recalculates the workbook in LibreOffice and compares it with the Python reference model.
Exits non-zero on any mismatch or formula error."""
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

import model

SRC = Path(__file__).resolve().parents[1] / "deliverables" / "Luca_Step1_Financial_Model.xlsx"
MC = [chr(ord("C") + i) for i in range(12)]
CHECKS = [  # (sheet, row, model key)
    ("P&L", 7, "revenue"), ("P&L", 8, "gm"), ("P&L", 18, None), ("P&L", 20, "op_profit"),
    ("P&L", 23, "np_nl"), ("P&L", 25, "np_wl"), ("Headcount", 14, "hc_total"), ("Headcount", 20, None),
    ("Cash Flow", 17, "cash_nl"), ("Cash Flow", 31, "cash_wl"), ("Runway", 9, "cash_nl"), ("Runway", 17, "cash_wl"),
]


def main():
    ref = model.run()
    ref_opex = [-x for x in ref["opex"]]
    ref_people = [a + b + c for a, b, c in zip(ref["salaries"], ref["benefits"], ref["hiring_costs"])]
    tmp = tempfile.mkdtemp()
    try:
        subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, str(SRC)],
                       check=True, capture_output=True)
        wb = load_workbook(Path(tmp) / SRC.name, data_only=True)
    finally:
        pass
    errors = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and re.match(r"^(#[A-Z/0-9]+[!?]|#N/A|Err:)", c.value):
                    print(f"formula error {ws.title}!{c.coordinate}: {c.value}")
                    errors += 1
    n = 0
    for sheet, row, key in CHECKS:
        expected = ref[key] if key else (ref_opex if sheet == "P&L" else ref_people)
        for m, col in enumerate(MC):
            got = wb[sheet][f"{col}{row}"].value
            n += 1
            if got is None or abs(got - expected[m]) > 0.5:
                print(f"mismatch {sheet}!{col}{row}: sheet={got} model={expected[m]}")
                errors += 1
    hires = wb["Headcount"]["D22"].value
    print(f"{n} values compared, {errors} problems; 25-hire check = {hires}; "
          f"runway = {wb['Runway']['C30'].value} / {wb['Runway']['D30'].value}; "
          f"cash M12 = {wb['Runway']['C24'].value:,.0f} / {wb['Runway']['D24'].value:,.0f}")
    shutil.rmtree(tmp)
    sys.exit(1 if errors or hires != "OK" else 0)


if __name__ == "__main__":
    main()
